import io
from datetime import date
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from .models import DietMenu, DietPlan, Menu


class MealPlanBulkUploadTestCase(TestCase):
    def setUp(self):
        self.upload_url = reverse("meals:mealplan_bulk_upload")
        self.list_url = reverse("meals:mealplan_list")

    def test_csv_bulk_upload_success(self):
        csv_content = (
            "date,meal_type,menu_name,category\n"
            "2026-05-20,중식,제육볶음,주반찬\n"
            "2026-05-20,중식,미역국,국\n"
            "2026-05-20,석식,카레라이스,일품\n"
        ).encode("utf-8-sig")

        file_obj = SimpleUploadedFile("test_meal.csv", csv_content, content_type="text/csv")
        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})

        self.assertRedirects(response, self.list_url)

        # Check DietPlans created
        lunch_plan = DietPlan.objects.get(target_date=date(2026, 5, 20), meal_type="중식")
        dinner_plan = DietPlan.objects.get(target_date=date(2026, 5, 20), meal_type="석식")

        lunch_menus = list(lunch_plan.diet_menus.values_list("menu__name", flat=True))
        self.assertIn("제육볶음", lunch_menus)
        self.assertIn("미역국", lunch_menus)

        dinner_menus = list(dinner_plan.diet_menus.values_list("menu__name", flat=True))
        self.assertIn("카레라이스", dinner_menus)

    def test_xlsx_bulk_upload_success(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["date", "meal_type", "menu_name", "category"])
        ws.append(["2026-05-21", "조식", "전복죽", "죽"])
        ws.append(["2026-05-21", "조식", "계란말이", "부반찬"])

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)

        file_obj = SimpleUploadedFile(
            "test_meal.xlsx",
            file_stream.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})
        self.assertRedirects(response, self.list_url)

        breakfast_plan = DietPlan.objects.get(target_date=date(2026, 5, 21), meal_type="조식")
        breakfast_menus = list(breakfast_plan.diet_menus.values_list("menu__name", flat=True))
        self.assertIn("전복죽", breakfast_menus)
        self.assertIn("계란말이", breakfast_menus)

    def test_xltx_bulk_upload_success(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["date", "meal_type", "menu_name", "category"])
        ws.append(["2026-05-22", "석식", "돈까스", "주반찬"])

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)

        file_obj = SimpleUploadedFile(
            "template_meal.xltx",
            file_stream.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.template",
        )

        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})
        self.assertRedirects(response, self.list_url)

        dinner_plan = DietPlan.objects.get(target_date=date(2026, 5, 22), meal_type="석식")
        dinner_menus = list(dinner_plan.diet_menus.values_list("menu__name", flat=True))
        self.assertIn("돈까스", dinner_menus)

    def test_grid_calendar_excel_upload_success(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["구분", "08-10(월)", "08-11(화)"])
        ws.append(["조식", "쌀밥", "쌀밥"])
        ws.append(["", "계란국", "콩나물국"])
        ws.append(["중식", "제육볶음", "돈까스"])
        ws.append(["", "미역국", "스프"])

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)

        file_obj = SimpleUploadedFile(
            "weekly_grid_meal.xlsx",
            file_stream.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})
        self.assertRedirects(response, self.list_url)

        current_year = date.today().year

        # Check August 10 Breakfast & Lunch
        aug10_bf = DietPlan.objects.get(target_date=date(current_year, 8, 10), meal_type="조식")
        bf_menus = list(aug10_bf.diet_menus.values_list("menu__name", flat=True))
        self.assertIn("쌀밥", bf_menus)
        self.assertIn("계란국", bf_menus)

        aug10_lunch = DietPlan.objects.get(target_date=date(current_year, 8, 10), meal_type="중식")
        lunch_menus = list(aug10_lunch.diet_menus.values_list("menu__name", flat=True))
        self.assertIn("제육볶음", lunch_menus)
        self.assertIn("미역국", lunch_menus)

    def test_grid_calendar_with_title_and_disclaimer_noise(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["2026년 8월 주간 식단표", None, None, None, None])
        ws.append(["구분", "08-10(월)", "08-11(화)", "08-12(수)", "08-13(목)"])
        ws.append(["조식", "쌀밥", "쌀밥", "잡곡밥", "쌀밥"])
        ws.append(["중식", "제육볶음", "돈까스", "카레라이스", "불고기"])
        ws.append(["석식", "볶음밥", "짜장면", "우동", "떡볶이"])
        ws.append(["※ 위 식단은 물품수급 사정에 따라 변경될 수 있습니다.", None, None, None, None])
        ws.append(["원산지 표시: 쌀(국산), 돼지고기(국산)", None, None, None, None])

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)

        file_obj = SimpleUploadedFile(
            "complex_grid_meal.xlsx",
            file_stream.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})
        self.assertRedirects(response, self.list_url)

        current_year = date.today().year

        # 1. Check all 4 date columns (10th, 11th, 12th, 13th) were extracted
        for day in (10, 11, 12, 13):
            plan = DietPlan.objects.get(target_date=date(current_year, 8, day), meal_type="중식")
            self.assertTrue(plan.diet_menus.exists())

        # 2. Check disclaimer text was NOT saved as menu
        all_menus = set(Menu.objects.values_list("name", flat=True))
        self.assertNotIn("※ 위 식단은 물품수급 사정에 따라 변경될 수 있습니다.", all_menus)
        self.assertNotIn("원산지 표시: 쌀(국산), 돼지고기(국산)", all_menus)

    def test_menu_sorting_by_category_priority(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["date", "meal_type", "menu_name", "category"])
        # Add menus in unsorted order: Fruit -> Kimchi -> Soup -> Main dish -> Rice -> Milk
        ws.append(["2026-05-25", "중식", "사과", "기타"])
        ws.append(["2026-05-25", "중식", "배추김치", "기타"])
        ws.append(["2026-05-25", "중식", "미역국", "기타"])
        ws.append(["2026-05-25", "중식", "제육볶음", "기타"])
        ws.append(["2026-05-25", "중식", "쌀밥", "기타"])
        ws.append(["2026-05-25", "중식", "우유", "기타"])

        file_stream = io.BytesIO()
        wb.save(file_stream)
        file_stream.seek(0)

        file_obj = SimpleUploadedFile(
            "unsorted_menu.xlsx",
            file_stream.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})
        self.assertRedirects(response, self.list_url)

        plan = DietPlan.objects.get(target_date=date(2026, 5, 25), meal_type="중식")
        saved_menus = list(
            DietMenu.objects.filter(diet_plan=plan)
            .order_by("id")
            .values_list("menu__name", flat=True)
        )

        # Expected priority order: 밥(1) -> 미역국(2) -> 제육볶음(3) -> 배추김치(4) -> 사과(5) -> 우유(6)
        expected_order = ["쌀밥", "미역국", "제육볶음", "배추김치", "사과", "우유"]
        self.assertEqual(saved_menus, expected_order)

    def test_bulk_upload_update_or_create_overwrite(self):
        # Create existing diet plan with menu
        plan = DietPlan.objects.create(target_date=date(2026, 5, 20), meal_type="중식")
        old_menu = Menu.objects.create(name="구식메뉴", category="주반찬")
        DietMenu.objects.create(diet_plan=plan, menu=old_menu)

        csv_content = (
            "date,meal_type,menu_name,category\n"
            "2026-05-20,중식,신식메뉴,주반찬\n"
        ).encode("utf-8-sig")

        file_obj = SimpleUploadedFile("update_meal.csv", csv_content, content_type="text/csv")
        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})

        self.assertRedirects(response, self.list_url)

        plan.refresh_from_db()
        menus = list(plan.diet_menus.values_list("menu__name", flat=True))
        self.assertEqual(menus, ["신식메뉴"])
        self.assertNotIn("구식메뉴", menus)

    def test_missing_required_headers(self):
        csv_content = "wrong_header,meal_type,menu_name\n2026-05-20,중식,제육볶음\n".encode("utf-8-sig")
        file_obj = SimpleUploadedFile("bad_header.csv", csv_content, content_type="text/csv")

        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})
        self.assertRedirects(response, self.upload_url)

    def test_invalid_date_format(self):
        csv_content = (
            "date,meal_type,menu_name,category\n"
            "invalid-date,중식,제육볶음,주반찬\n"
        ).encode("utf-8-sig")
        file_obj = SimpleUploadedFile("bad_date.csv", csv_content, content_type="text/csv")

        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})
        self.assertRedirects(response, self.upload_url)

    def test_invalid_meal_type(self):
        csv_content = (
            "date,meal_type,menu_name,category\n"
            "2026-05-20,야식,제육볶음,주반찬\n"
        ).encode("utf-8-sig")
        file_obj = SimpleUploadedFile("bad_mealtype.csv", csv_content, content_type="text/csv")

        response = self.client.post(self.upload_url, {"mealplan_file": file_obj})
        self.assertRedirects(response, self.upload_url)

